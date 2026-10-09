from datetime import datetime
from io import BytesIO
from typing import Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

PERIOD_LABELS = {
    "today": "Сегодня",
    "week": "Последние 7 дней",
    "month": "Последние 30 дней",
}
HEADERS = ["№", "ГРЗ", "Дата", "Время", "Адрес", "Статус"]

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def period_label(period) -> str:
    return PERIOD_LABELS.get(period, "Всё время")


def status_label(reaction_status, found_status) -> str:
    if found_status == "Да":
        return "Найден"
    if reaction_status == "Принято":
        return "Принято"
    if reaction_status == "Отклонено":
        return "Отклонено"
    return "Ожидает"


def _rows(alerts: Sequence) -> list:
    return [
        [i, a.vehicle_plate, a.alert_date, a.alert_time, a.address,
         status_label(a.reaction_status, a.found_status)]
        for i, a in enumerate(alerts, start=1)
    ]


def build_history_docx(alerts: Sequence, employee_name: str, period) -> bytes:
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("История оповещений")
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph(f"Сотрудник: {employee_name}")
    doc.add_paragraph(f"Период: {period_label(period)}")
    doc.add_paragraph(f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}")

    table = doc.add_table(rows=1, cols=len(HEADERS))
<<<<<<< HEAD
    table.style = "Table Grid"
=======
    table.style = "Table Grid"  
>>>>>>> cc988d8a95767965fb2bdaa7c8fe905d435e39b1

    for cell, text in zip(table.rows[0].cells, HEADERS):
        cell.paragraphs[0].add_run(text).bold = True

    for row in _rows(alerts):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)

    if not alerts:
        doc.add_paragraph("За выбранный период оповещений нет.")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_history_xlsx(alerts: Sequence, employee_name: str, period) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "История оповещений"

    ws["A1"] = "История оповещений"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Сотрудник: {employee_name}"
    ws["A3"] = f"Период: {period_label(period)}"
    ws["A4"] = f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}"

    header_row = 6
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, text in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(_rows(alerts), start=header_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = border

    for col, width in zip("ABCDEF", (5, 14, 12, 10, 48, 12)):
        ws.column_dimensions[col].width = width

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()